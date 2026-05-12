from __future__ import annotations

import json
import queue
import shutil
import threading
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import customtkinter as ctk
from docx import Document
from faster_whisper import WhisperModel
from tkinter import filedialog, messagebox

from video2text.downloader import download_audio


class Video2TextDesktop(ctk.CTk):
    def __init__(self) -> None:
        super().__init__()
        self.title("Video2Text Desktop")
        self.geometry("1200x780")
        self.minsize(1050, 700)

        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("green")

        self.selected_file: Path | None = None
        self.transcript_text = ""
        self.docx_bytes = b""
        self.log_queue: queue.Queue[tuple[str, Any]] = queue.Queue()

        self._build_ui()
        self.after(100, self._poll_queue)

    def _build_ui(self) -> None:
        self.configure(fg_color="#000000")
        self.grid_columnconfigure(0, weight=0)
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(1, weight=1)

        header = ctk.CTkFrame(self, fg_color="#000000", corner_radius=0)
        header.grid(row=0, column=0, columnspan=2, sticky="ew", padx=0, pady=0)
        header.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            header,
            text="VIDEO2TEXT // CYBER TRANSCRIBER",
            font=ctk.CTkFont(family="Consolas", size=28, weight="bold"),
            text_color="#00ff88",
        ).grid(row=0, column=0, sticky="w", padx=22, pady=(18, 2))

        ctk.CTkLabel(
            header,
            text="> app local de escritorio · carga video · pega URL · transcribe en tiempo real · exporta TXT y Word",
            font=ctk.CTkFont(family="Consolas", size=13),
            text_color="#b8ffd9",
        ).grid(row=1, column=0, sticky="w", padx=22, pady=(0, 14))

        sidebar = ctk.CTkScrollableFrame(
            self,
            width=320,
            fg_color="#021109",
            border_color="#00ff88",
            border_width=1,
            scrollbar_button_color="#063f27",
            scrollbar_button_hover_color="#00ff88",
        )
        sidebar.grid(row=1, column=0, sticky="nsw", padx=16, pady=12)
        sidebar.grid_propagate(False)

        ctk.CTkLabel(
            sidebar,
            text="PANEL DE CONTROL",
            font=ctk.CTkFont(family="Consolas", size=18, weight="bold"),
            text_color="#00ff88",
        ).pack(anchor="w", padx=16, pady=(14, 10))

        self.file_label = ctk.CTkLabel(
            sidebar,
            text="Archivo: ninguno",
            wraplength=260,
            justify="left",
            font=ctk.CTkFont(family="Consolas", size=12),
            text_color="#b8ffd9",
        )
        self.file_label.pack(anchor="w", padx=16, pady=(0, 6))

        ctk.CTkButton(
            sidebar,
            text="Cargar Archivo (video)",
            command=self._select_file,
            fg_color="#052b1b",
            hover_color="#0b6b42",
            text_color="#00ff88",
            border_color="#00ff88",
            border_width=1,
        ).pack(fill="x", padx=16, pady=(4, 12))

        ctk.CTkLabel(
            sidebar,
            text="Pegar URL",
            font=ctk.CTkFont(family="Consolas", size=13, weight="bold"),
            text_color="#00ff88",
        ).pack(anchor="w", padx=16, pady=(0, 4))

        self.url_entry = ctk.CTkEntry(
            sidebar,
            placeholder_text="https://www.youtube.com/watch?v=...",
            font=ctk.CTkFont(family="Consolas", size=12),
            fg_color="#000000",
            border_color="#00ff88",
            text_color="#b8ffd9",
            height=36,
        )
        self.url_entry.pack(fill="x", padx=16, pady=(0, 10))

        self.process_button = ctk.CTkButton(
            sidebar,
            text="▶ PROCESAR VIDEO",
            height=46,
            command=self._start_processing,
            fg_color="#00ff88",
            hover_color="#00cc70",
            text_color="#001b0d",
            font=ctk.CTkFont(family="Consolas", size=16, weight="bold"),
        )
        self.process_button.pack(fill="x", padx=16, pady=(4, 8))

        self.download_txt_button = ctk.CTkButton(
            sidebar,
            text="Descargar TXT",
            command=self._save_txt,
            state="disabled",
            fg_color="#052b1b",
            hover_color="#0b6b42",
            text_color="#00ff88",
            border_color="#00ff88",
            border_width=1,
        )
        self.download_txt_button.pack(fill="x", padx=16, pady=(2, 6))

        self.download_word_button = ctk.CTkButton(
            sidebar,
            text="Descargar Word",
            command=self._save_docx,
            state="disabled",
            fg_color="#052b1b",
            hover_color="#0b6b42",
            text_color="#00ff88",
            border_color="#00ff88",
            border_width=1,
        )
        self.download_word_button.pack(fill="x", padx=16, pady=(0, 12))

        self.model_option = self._option(sidebar, "Modelo", ["tiny", "base", "small", "medium", "large-v3"], "small")
        self.language_option = self._option(sidebar, "Idioma", ["auto", "es", "en", "pt", "fr", "de", "it"], "auto")
        self.device_option = self._option(sidebar, "Dispositivo", ["cpu", "cuda"], "cpu")
        self.compute_option = self._option(sidebar, "Compute type", ["int8", "float32", "float16"], "int8")
        self.cookies_option = self._option(sidebar, "Cookies navegador", ["No usar", "chrome", "edge", "firefox"], "No usar")

        self.status_label = ctk.CTkLabel(
            sidebar,
            text="Estado: listo",
            wraplength=260,
            justify="left",
            font=ctk.CTkFont(family="Consolas", size=12),
            text_color="#b8ffd9",
        )
        self.status_label.pack(anchor="w", padx=16, pady=(12, 18))

        main_panel = ctk.CTkFrame(self, fg_color="#000000", border_color="#00ff88", border_width=1)
        main_panel.grid(row=1, column=1, sticky="nsew", padx=(0, 16), pady=12)
        main_panel.grid_columnconfigure(0, weight=1)
        main_panel.grid_rowconfigure(1, weight=1)

        ctk.CTkLabel(
            main_panel,
            text="PANTALLA DE TRANSCRIPCIÓN EN TIEMPO REAL",
            font=ctk.CTkFont(family="Consolas", size=18, weight="bold"),
            text_color="#00ff88",
        ).grid(row=0, column=0, sticky="w", padx=18, pady=(18, 8))

        self.terminal = ctk.CTkTextbox(
            main_panel,
            font=ctk.CTkFont(family="Consolas", size=13),
            fg_color="#000000",
            text_color="#00ff88",
            border_color="#00ff88",
            border_width=1,
            wrap="word",
        )
        self.terminal.grid(row=1, column=0, sticky="nsew", padx=18, pady=(0, 12))
        self.terminal.insert("end", "[READY] Esperando archivo o URL...\n[INFO] Todo se procesa localmente en este equipo.\n")
        self.terminal.configure(state="disabled")

        self.progress = ctk.CTkProgressBar(main_panel, progress_color="#00ff88")
        self.progress.grid(row=2, column=0, sticky="ew", padx=18, pady=(0, 18))
        self.progress.set(0)

    def _option(self, parent: ctk.CTkScrollableFrame, label: str, values: list[str], default: str) -> ctk.CTkOptionMenu:
        ctk.CTkLabel(
            parent,
            text=label,
            font=ctk.CTkFont(family="Consolas", size=13),
            text_color="#00ff88",
        ).pack(anchor="w", padx=16, pady=(8, 4))
        option = ctk.CTkOptionMenu(
            parent,
            values=values,
            fg_color="#000000",
            button_color="#063f27",
            button_hover_color="#0b6b42",
            text_color="#b8ffd9",
        )
        option.set(default)
        option.pack(fill="x", padx=16, pady=(0, 4))
        return option

    def _select_file(self) -> None:
        file_path = filedialog.askopenfilename(
            title="Seleccionar video o audio",
            filetypes=[
                ("Video/audio", "*.mp4 *.mov *.mkv *.webm *.mp3 *.wav *.m4a *.aac *.ogg"),
                ("Todos los archivos", "*.*"),
            ],
        )
        if file_path:
            self.selected_file = Path(file_path)
            self.file_label.configure(text=f"Archivo: {self.selected_file.name}")

    def _start_processing(self) -> None:
        url = self.url_entry.get().strip()
        if not self.selected_file and not url:
            messagebox.showwarning("Falta fuente", "Carga un archivo o pega una URL.")
            return
        self._set_busy(True)
        self._clear_terminal()
        self.progress.set(0)
        self.transcript_text = ""
        self.docx_bytes = b""
        self.download_txt_button.configure(state="disabled")
        self.download_word_button.configure(state="disabled")
        self.status_label.configure(text="Estado: procesando...")
        threading.Thread(target=self._process_worker, daemon=True).start()

    def _process_worker(self) -> None:
        try:
            base_output = Path("outputs") / datetime.now().strftime("desktop_%Y%m%d_%H%M%S")
            base_output.mkdir(parents=True, exist_ok=True)
            self.log_queue.put(("log", "[BOOT] Iniciando Video2Text Desktop..."))

            media_path: Path
            metadata: dict[str, Any]
            url = self.url_entry.get().strip()

            if self.selected_file:
                self.log_queue.put(("log", "[INPUT] Archivo local seleccionado."))
                uploads_dir = base_output / "uploaded"
                uploads_dir.mkdir(parents=True, exist_ok=True)
                media_path = uploads_dir / self.selected_file.name
                shutil.copy2(self.selected_file, media_path)
                metadata = {"title": self.selected_file.name, "source": "local_file"}
            else:
                self.log_queue.put(("log", "[INPUT] URL recibida. Descargando con yt-dlp..."))
                cookies_value = self.cookies_option.get()
                cookies = None if cookies_value == "No usar" else cookies_value
                media_path, metadata = download_audio(url, base_output / "downloads", cookies_from_browser=cookies)
                self._write_metadata(metadata, base_output)

            model_size = self.model_option.get()
            language_value = self.language_option.get()
            language = None if language_value == "auto" else language_value
            device = self.device_option.get()
            compute_type = self.compute_option.get()

            self.log_queue.put(("log", f"[MEDIA] Archivo preparado: {media_path.name}"))
            self.log_queue.put(("log", f"[MODEL] Cargando modelo: {model_size}"))

            model = WhisperModel(model_size, device=device, compute_type=compute_type)
            segments_iter, info = model.transcribe(str(media_path), language=language, vad_filter=True, beam_size=5)

            duration = float(info.duration or 0)
            self.log_queue.put(("log", f"[LANG] Idioma detectado: {info.language} ({info.language_probability:.2%})"))
            self.log_queue.put(("log", "[LIVE] Transcripción en tiempo real:"))
            self.log_queue.put(("log", ""))

            segments: list[dict[str, Any]] = []
            parts: list[str] = []
            for segment in segments_iter:
                item = {"id": segment.id, "start": segment.start, "end": segment.end, "text": segment.text.strip()}
                segments.append(item)
                parts.append(item["text"])
                self.log_queue.put(("log", f"[{self._stamp(segment.start)} --> {self._stamp(segment.end)}] {item['text']}"))
                if duration > 0:
                    self.log_queue.put(("progress", min(float(segment.end) / duration, 1.0)))

            transcript = "\n".join(parts).strip() + "\n"
            (base_output / "transcript.txt").write_text(transcript, encoding="utf-8")
            (base_output / "segments.json").write_text(json.dumps(segments, ensure_ascii=False, indent=2), encoding="utf-8")
            self._write_srt(segments, base_output / "subtitles.srt")
            docx_bytes = self._make_docx(transcript, metadata)
            (base_output / "transcript.docx").write_bytes(docx_bytes)

            self.log_queue.put(("done", {"transcript": transcript, "docx": docx_bytes, "output": str(base_output)}))
        except Exception as exc:
            self.log_queue.put(("error", str(exc)))

    def _poll_queue(self) -> None:
        try:
            while True:
                msg_type, payload = self.log_queue.get_nowait()
                if msg_type == "log":
                    self._log(str(payload))
                elif msg_type == "progress":
                    self.progress.set(float(payload))
                elif msg_type == "done":
                    self.transcript_text = payload["transcript"]
                    self.docx_bytes = payload["docx"]
                    self.progress.set(1)
                    self.download_txt_button.configure(state="normal")
                    self.download_word_button.configure(state="normal")
                    self._set_busy(False)
                    self.status_label.configure(text=f"Estado: listo. Salida: {payload['output']}")
                    self._log("\n[DONE] Archivos generados. Ya puedes descargar TXT o Word.")
                elif msg_type == "error":
                    self._set_busy(False)
                    self.status_label.configure(text="Estado: error")
                    self._log(f"[ERROR] {payload}")
                    messagebox.showerror("Error", str(payload))
        except queue.Empty:
            pass
        self.after(100, self._poll_queue)

    def _set_busy(self, busy: bool) -> None:
        self.process_button.configure(state="disabled" if busy else "normal")

    def _log(self, text: str) -> None:
        self.terminal.configure(state="normal")
        self.terminal.insert("end", text + "\n")
        self.terminal.see("end")
        self.terminal.configure(state="disabled")

    def _clear_terminal(self) -> None:
        self.terminal.configure(state="normal")
        self.terminal.delete("1.0", "end")
        self.terminal.configure(state="disabled")

    def _save_txt(self) -> None:
        path = filedialog.asksaveasfilename(title="Guardar TXT", defaultextension=".txt", filetypes=[("Texto", "*.txt")], initialfile="transcripcion.txt")
        if path:
            Path(path).write_text(self.transcript_text, encoding="utf-8")
            messagebox.showinfo("Guardado", "TXT guardado correctamente.")

    def _save_docx(self) -> None:
        path = filedialog.asksaveasfilename(title="Guardar Word", defaultextension=".docx", filetypes=[("Word", "*.docx")], initialfile="transcripcion.docx")
        if path:
            Path(path).write_bytes(self.docx_bytes)
            messagebox.showinfo("Guardado", "Word guardado correctamente.")

    @staticmethod
    def _stamp(seconds: float) -> str:
        milliseconds = int(round(float(seconds) * 1000))
        hours, remainder = divmod(milliseconds, 3_600_000)
        minutes, remainder = divmod(remainder, 60_000)
        secs, millis = divmod(remainder, 1000)
        return f"{hours:02}:{minutes:02}:{secs:02}.{millis:03}"

    @classmethod
    def _write_srt(cls, segments: list[dict[str, Any]], path: Path) -> None:
        lines: list[str] = []
        for index, segment in enumerate(segments, start=1):
            start = cls._stamp(segment["start"]).replace(".", ",")
            end = cls._stamp(segment["end"]).replace(".", ",")
            lines.extend([str(index), f"{start} --> {end}", str(segment["text"]).strip(), ""])
        path.write_text("\n".join(lines), encoding="utf-8")

    @staticmethod
    def _write_metadata(metadata: dict[str, Any], output_dir: Path) -> None:
        keep = ["id", "title", "webpage_url", "extractor", "duration", "uploader", "upload_date", "description"]
        clean = {key: metadata.get(key) for key in keep if key in metadata}
        (output_dir / "metadata.json").write_text(json.dumps(clean, ensure_ascii=False, indent=2), encoding="utf-8")

    @staticmethod
    def _make_docx(transcript: str, metadata: dict[str, Any] | None = None) -> bytes:
        document = Document()
        document.add_heading("Video2Text - Transcripción", level=1)
        document.add_paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        if metadata:
            document.add_heading("Metadatos", level=2)
            for key in ["title", "webpage_url", "uploader", "duration", "extractor", "source"]:
                if metadata.get(key):
                    document.add_paragraph(f"{key}: {metadata.get(key)}")
        document.add_heading("Transcripción", level=2)
        for paragraph in transcript.split("\n"):
            paragraph = paragraph.strip()
            if paragraph:
                document.add_paragraph(paragraph)
        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


def main() -> None:
    app = Video2TextDesktop()
    app.mainloop()


if __name__ == "__main__":
    main()
