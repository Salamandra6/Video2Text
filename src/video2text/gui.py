from __future__ import annotations

import json
import shutil
import tempfile
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import streamlit as st
from docx import Document
from faster_whisper import WhisperModel

from video2text.downloader import download_audio


APP_TITLE = "Video2Text // Cyber Transcriber"


def inject_hacker_css() -> None:
    st.markdown(
        """
        <style>
        :root {
            --neon-green: #00ff88;
            --dark-bg: #020403;
            --panel-bg: rgba(0, 20, 12, 0.88);
            --text-soft: #b8ffd9;
        }

        .stApp {
            background:
                radial-gradient(circle at top left, rgba(0, 255, 136, 0.16), transparent 30%),
                radial-gradient(circle at bottom right, rgba(0, 140, 255, 0.12), transparent 28%),
                linear-gradient(135deg, #010201 0%, #03140c 45%, #000000 100%);
            color: var(--text-soft);
        }

        .block-container {
            padding-top: 2rem;
            max-width: 1180px;
        }

        h1, h2, h3 {
            color: var(--neon-green) !important;
            text-shadow: 0 0 14px rgba(0, 255, 136, 0.65);
            letter-spacing: 0.04em;
        }

        .cyber-panel {
            border: 1px solid rgba(0, 255, 136, 0.45);
            background: var(--panel-bg);
            border-radius: 18px;
            padding: 1.2rem;
            box-shadow: 0 0 24px rgba(0, 255, 136, 0.18);
        }

        .cyber-title {
            font-size: 2.4rem;
            font-weight: 900;
            color: var(--neon-green);
            text-shadow: 0 0 16px rgba(0, 255, 136, 0.85);
            margin-bottom: 0.2rem;
        }

        .cyber-subtitle {
            color: #b8ffd9;
            font-family: Consolas, monospace;
            margin-bottom: 1.4rem;
        }

        .terminal-box {
            height: 420px;
            overflow-y: auto;
            white-space: pre-wrap;
            font-family: Consolas, 'Courier New', monospace;
            color: #00ff88;
            background: rgba(0, 0, 0, 0.78);
            border: 1px solid rgba(0, 255, 136, 0.45);
            border-radius: 14px;
            padding: 1rem;
            box-shadow: inset 0 0 18px rgba(0, 255, 136, 0.14);
        }

        .status-line {
            font-family: Consolas, monospace;
            color: #00ff88;
        }

        .stButton > button, .stDownloadButton > button {
            border: 1px solid #00ff88 !important;
            background: rgba(0, 255, 136, 0.08) !important;
            color: #00ff88 !important;
            border-radius: 12px !important;
            font-weight: 800 !important;
            box-shadow: 0 0 14px rgba(0, 255, 136, 0.18);
        }

        .stButton > button:hover, .stDownloadButton > button:hover {
            background: rgba(0, 255, 136, 0.18) !important;
            box-shadow: 0 0 24px rgba(0, 255, 136, 0.38);
        }

        input, textarea {
            font-family: Consolas, monospace !important;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def safe_name(value: str) -> str:
    clean = "".join(ch if ch.isalnum() or ch in "._-" else "_" for ch in value)
    return clean[:90].strip("._-") or "video2text"


def make_docx(transcript: str, metadata: dict[str, Any] | None = None) -> bytes:
    document = Document()
    document.add_heading("Video2Text - Transcripción", level=1)
    document.add_paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    if metadata:
        document.add_heading("Metadatos", level=2)
        for key in ["title", "webpage_url", "uploader", "duration", "extractor"]:
            if metadata.get(key):
                document.add_paragraph(f"{key}: {metadata.get(key)}")

    document.add_heading("Transcripción", level=2)
    for paragraph in transcript.split("\n"):
        paragraph = paragraph.strip()
        if paragraph:
            document.add_paragraph(paragraph)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def save_metadata(metadata: dict[str, Any], output_dir: Path) -> None:
    keep_keys = [
        "id",
        "title",
        "webpage_url",
        "extractor",
        "duration",
        "uploader",
        "upload_date",
        "description",
    ]
    clean_metadata = {key: metadata.get(key) for key in keep_keys if key in metadata}
    (output_dir / "metadata.json").write_text(
        json.dumps(clean_metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def format_srt_timestamp(seconds: float) -> str:
    milliseconds = int(round(seconds * 1000))
    hours, remainder = divmod(milliseconds, 3_600_000)
    minutes, remainder = divmod(remainder, 60_000)
    secs, millis = divmod(remainder, 1000)
    return f"{hours:02}:{minutes:02}:{secs:02},{millis:03}"


def write_srt(segments: list[dict[str, Any]], output_path: Path) -> None:
    lines: list[str] = []
    for index, segment in enumerate(segments, start=1):
        start = format_srt_timestamp(float(segment["start"]))
        end = format_srt_timestamp(float(segment["end"]))
        text = str(segment["text"]).strip()
        lines.extend([str(index), f"{start} --> {end}", text, ""])
    output_path.write_text("\n".join(lines), encoding="utf-8")


def transcribe_realtime(
    media_path: Path,
    output_dir: Path,
    terminal_placeholder: Any,
    progress_bar: Any,
    model_size: str,
    language: str | None,
    device: str,
    compute_type: str,
) -> tuple[str, list[dict[str, Any]], dict[str, Any]]:
    output_dir.mkdir(parents=True, exist_ok=True)

    terminal_lines: list[str] = []

    def log(line: str) -> None:
        terminal_lines.append(line)
        terminal_placeholder.markdown(
            f"<div class='terminal-box'>{'\n'.join(terminal_lines)}</div>",
            unsafe_allow_html=True,
        )

    log("[BOOT] Inicializando motor de transcripción local...")
    log(f"[MODEL] Cargando modelo Whisper: {model_size}")
    model = WhisperModel(model_size, device=device, compute_type=compute_type)

    log("[SCAN] Analizando audio/video...")
    segments_iter, info = model.transcribe(
        str(media_path),
        language=language,
        vad_filter=True,
        beam_size=5,
    )

    segments: list[dict[str, Any]] = []
    transcript_parts: list[str] = []
    duration = float(info.duration or 0)

    log(f"[LANG] Idioma detectado: {info.language} ({info.language_probability:.2%})")
    log("[LIVE] Generando transcripción en tiempo real...")
    log("")

    for segment in segments_iter:
        item = {
            "id": segment.id,
            "start": segment.start,
            "end": segment.end,
            "text": segment.text.strip(),
        }
        segments.append(item)
        transcript_parts.append(item["text"])

        start = format_srt_timestamp(float(segment.start)).replace(",", ".")
        end = format_srt_timestamp(float(segment.end)).replace(",", ".")
        log(f"[{start} --> {end}] {item['text']}")

        if duration > 0:
            progress_bar.progress(min(float(segment.end) / duration, 1.0))

    transcript = "\n".join(transcript_parts).strip() + "\n"

    txt_path = output_dir / "transcript.txt"
    json_path = output_dir / "segments.json"
    srt_path = output_dir / "subtitles.srt"
    docx_path = output_dir / "transcript.docx"

    txt_path.write_text(transcript, encoding="utf-8")
    json_path.write_text(json.dumps(segments, ensure_ascii=False, indent=2), encoding="utf-8")
    write_srt(segments, srt_path)
    docx_path.write_bytes(make_docx(transcript))

    progress_bar.progress(1.0)
    log("")
    log("[DONE] Transcripción finalizada. Archivos listos para descarga.")

    result = {
        "language": info.language,
        "language_probability": info.language_probability,
        "duration": info.duration,
        "txt_path": str(txt_path),
        "json_path": str(json_path),
        "srt_path": str(srt_path),
        "docx_path": str(docx_path),
        "segments": len(segments),
    }
    return transcript, segments, result


def main() -> None:
    st.set_page_config(page_title=APP_TITLE, page_icon="🟢", layout="wide")
    inject_hacker_css()

    st.markdown("<div class='cyber-title'>Video2Text</div>", unsafe_allow_html=True)
    st.markdown(
        "<div class='cyber-subtitle'>> Sistema local de extracción y transcripción // modo hacker cinematográfico</div>",
        unsafe_allow_html=True,
    )

    if "transcript" not in st.session_state:
        st.session_state.transcript = ""
    if "docx_bytes" not in st.session_state:
        st.session_state.docx_bytes = b""
    if "metadata" not in st.session_state:
        st.session_state.metadata = {}

    left, right = st.columns([0.38, 0.62])

    with left:
        st.markdown("<div class='cyber-panel'>", unsafe_allow_html=True)
        st.subheader("Panel de control")

        uploaded_file = st.file_uploader(
            "Cargar Archivo (video)",
            type=["mp4", "mov", "mkv", "webm", "mp3", "wav", "m4a", "aac", "ogg"],
        )

        url = st.text_input("Pegar URL", placeholder="https://www.youtube.com/watch?v=...")

        model_size = st.selectbox(
            "Modelo",
            ["tiny", "base", "small", "medium", "large-v3"],
            index=2,
        )
        language_value = st.selectbox(
            "Idioma",
            ["auto", "es", "en", "pt", "fr", "de", "it"],
            index=0,
        )
        device = st.selectbox("Dispositivo", ["cpu", "cuda"], index=0)
        compute_type = st.selectbox("Compute type", ["int8", "float32", "float16"], index=0)
        cookies_browser = st.selectbox("Cookies del navegador", ["No usar", "chrome", "edge", "firefox"], index=0)

        process = st.button("Procesar Video", use_container_width=True)

        st.divider()

        st.download_button(
            "Descargar TXT",
            data=st.session_state.transcript.encode("utf-8") if st.session_state.transcript else b"",
            file_name="transcripcion.txt",
            mime="text/plain",
            disabled=not bool(st.session_state.transcript),
            use_container_width=True,
        )
        st.download_button(
            "Descargar Word",
            data=st.session_state.docx_bytes if st.session_state.docx_bytes else b"",
            file_name="transcripcion.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            disabled=not bool(st.session_state.docx_bytes),
            use_container_width=True,
        )
        st.markdown("</div>", unsafe_allow_html=True)

    with right:
        st.subheader("Pantalla de transcripción en tiempo real")
        terminal_placeholder = st.empty()
        progress_bar = st.progress(0)
        terminal_placeholder.markdown(
            "<div class='terminal-box'>[READY] Esperando archivo o URL...\n[INFO] Todo se procesa localmente en este equipo.</div>",
            unsafe_allow_html=True,
        )

    if process:
        if not uploaded_file and not url.strip():
            st.error("Debes cargar un archivo o pegar una URL.")
            return

        base_output = Path("outputs") / safe_name(datetime.now().strftime("gui_%Y%m%d_%H%M%S"))
        base_output.mkdir(parents=True, exist_ok=True)
        media_path: Path
        metadata: dict[str, Any] = {}

        with st.spinner("Preparando fuente del video..."):
            if uploaded_file:
                uploads_dir = base_output / "uploaded"
                uploads_dir.mkdir(parents=True, exist_ok=True)
                media_path = uploads_dir / safe_name(uploaded_file.name)
                media_path.write_bytes(uploaded_file.getbuffer())
                metadata = {"title": uploaded_file.name, "source": "uploaded_file"}
            else:
                downloads_dir = base_output / "downloads"
                cookies = None if cookies_browser == "No usar" else cookies_browser
                media_path, metadata = download_audio(url.strip(), downloads_dir, cookies_from_browser=cookies)
                save_metadata(metadata, base_output)

        language = None if language_value == "auto" else language_value

        transcript, _segments, result = transcribe_realtime(
            media_path=media_path,
            output_dir=base_output,
            terminal_placeholder=terminal_placeholder,
            progress_bar=progress_bar,
            model_size=model_size,
            language=language,
            device=device,
            compute_type=compute_type,
        )

        docx_bytes = make_docx(transcript, metadata)
        docx_path = Path(result["docx_path"])
        docx_path.write_bytes(docx_bytes)

        st.session_state.transcript = transcript
        st.session_state.docx_bytes = docx_bytes
        st.session_state.metadata = metadata

        st.success("Transcripción completada. Ya puedes descargar TXT o Word.")


if __name__ == "__main__":
    main()
